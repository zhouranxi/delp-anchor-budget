"""Generate the Elsevier declaration files for the S2 submission.

PR (and Elsevier generally) requires the Declaration of Competing Interests to
be produced with the Declarations Tool (https://declarations.elsevier.com/),
which needs an e-mail one-time passcode and therefore cannot be scripted.

This script produces a ready-to-upload .docx with the exact Elsevier-standard
wording, so that if the tool is unavailable the author can upload this
instead -- and it doubles as a checklist of what to type into the tool.

Run from the package root:  python audit/make_declarations_docx.py
"""
import os

from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "declarations.docx")

TITLE = "The Anchor Accumulation Problem: Bounded Anchor Sets Recover " \
        "Plasticity in Usage-Localized Elastic Weight Consolidation"
AUTHOR = "Ranxi Zhou"
JOURNAL = "Pattern Recognition"

# --- Elsevier-standard wording -------------------------------------------
COMPETING = (
    "The author declares that he has no known competing financial interests "
    "or personal relationships that could have appeared to influence the "
    "work reported in this paper."
)

FUNDING = (
    "This research did not receive any specific grant from funding agencies "
    "in the public, commercial, or not-for-profit sectors."
)

DATA = (
    "All code, experiment scripts, and the complete JSON records of every run "
    "(including per-phase accuracy matrices and anchor-stack survivor "
    "histories) are publicly available at "
    "https://github.com/zhouranxi/delp-anchor-budget. This manuscript and its "
    "full reproducibility package -- including the audit scripts that "
    "re-derive every quoted number from those records -- are archived on "
    "Zenodo (doi:10.5281/zenodo.23668522). The precursor audit study is "
    "archived at doi:10.5281/zenodo.23572226 and the companion "
    "segment-specialization audit at doi:10.5281/zenodo.23637972; the "
    "underlying dataset and code release is doi:10.5281/zenodo.23639561."
)

CRediT = (
    "Ranxi Zhou: Conceptualization, Methodology, Software, Investigation, "
    "Formal analysis, Data curation, Writing - original draft, "
    "Writing - review & editing, Visualization."
)

PREPRINT = (
    "A precursor version of this work by the same author is archived as a "
    "preprint on Zenodo (doi:10.5281/zenodo.23572226) and is cited and "
    "clearly distinguished from the present work. The present manuscript is "
    "original, has not been published previously, and is not under "
    "consideration by any other journal."
)

SECTIONS = [
    ("Declaration of competing interest", COMPETING),
    ("Funding", FUNDING),
    ("Data availability", DATA),
    ("CRediT authorship contribution statement", CRediT),
    ("Preprint and prior work", PREPRINT),
]


def build():
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)

    doc.add_heading("Declarations", level=1)

    p = doc.add_paragraph()
    p.add_run("Manuscript title: ").bold = True
    p.add_run(TITLE)
    p = doc.add_paragraph()
    p.add_run("Corresponding author: ").bold = True
    p.add_run(f"{AUTHOR} (Hunan University of Chinese Medicine, China)")
    p = doc.add_paragraph()
    p.add_run("Target journal: ").bold = True
    p.add_run(JOURNAL)

    doc.add_paragraph()

    for heading, body in SECTIONS:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(body)
        doc.add_paragraph()

    doc.save(OUT)
    print(f"written: {OUT}")
    print(f"size   : {os.path.getsize(OUT)} bytes")
    print()
    print("Reminder: PR asks for this file to be produced by the Declarations")
    print("Tool (https://declarations.elsevier.com/). If you use the tool,")
    print("upload the tool's own .docx instead of this file.")


if __name__ == "__main__":
    build()
